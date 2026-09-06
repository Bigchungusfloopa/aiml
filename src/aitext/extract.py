"""Plain-text extraction from uploaded files (PDF / DOCX / TXT).

Used by the web UI so a user can drop a document instead of pasting text.
Everything is best-effort and never raises to the caller.
"""

from __future__ import annotations

import io

SUPPORTED = ("txt", "text", "md", "pdf", "docx")


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> tuple[str, str | None]:
    """Return ``(text, error)``. ``error`` is None on success."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if ext in ("txt", "text", "md", ""):
            return data.decode("utf-8", errors="replace"), None
        if ext == "pdf":
            text = _from_pdf(data)
            if not text.strip():
                return "", ("No selectable text found in the PDF "
                            "(it may be a scanned image).")
            return text, None
        if ext == "docx":
            return _from_docx(data), None
        return "", f"Unsupported file type: .{ext}"
    except Exception as exc:  # noqa: BLE001
        return "", f"Could not read {filename}: {exc}"
